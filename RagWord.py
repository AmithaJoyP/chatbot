import os
import groq
DOC_PATH = os.path.join(os.path.dirname(__file__), "RAG", "word.docx")
print(os.path.dirname(__file__))
print(os.path.join(os.path.dirname(__file__), "data", "word.docx"))
from docx import Document 
doc = Document("word.docx")
for paragraph in doc.paragraphs:
    print(paragraph.text)
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.embeddings import SentenceTransformerEmbeddings
from langchain_community.vectorstores import Chroma
from groq import Groq
from langchain.prompts import PromptTemplate
from langchain.schema.runnable import RunnablePassthrough
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma
from langchain.schema.output_parser import StrOutputParser




GROQ_API_KEY = os.getenv("GROQ_API_KEY") 
groq_client = groq.Groq(api_key=GROQ_API_KEY)
print(groq_client.models.list())

# Configuration
DOC_PATH = "C:/Users/admin/Documents/VS CODE/RAG/word.docx"
CHROMA_DB_DIR = "chroma_db_word_rag"
EMBEDDING_MODEL_NAME = "all-MiniLM-L6-v2"

def load_and_chunk_docx(file_path):
    try:
        doc = Document(file_path)
        full_text = [paragraph.text for paragraph in doc.paragraphs]
        text = "\n".join(full_text)
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000,chunk_overlap=200,length_function=len,is_separator_regex=False)
        chunks = text_splitter.create_documents([text])
        print(f"Loaded {len(chunks)} chunks from '{file_path}'")
        return chunks
    except Exception as e:
        print(f"Error loading or chunking document: {e}")
        return []

def initialize_chroma_db(chunks, persist_directory):
    try:
        embeddings = SentenceTransformerEmbeddings(model_name=EMBEDDING_MODEL_NAME)
        print(f"Loading embedding model: {EMBEDDING_MODEL_NAME}...")
        if not os.path.exists(persist_directory):
            print(f"Creating new ChromaDB at '{persist_directory}'...")
            db = Chroma.from_documents(documents=chunks,embedding=embeddings,persist_directory=persist_directory)
            db.persist()
            print("ChromaDB created and persisted.")
        else:
            print(f"Loading existing ChromaDB from '{persist_directory}'...")
            db = Chroma(persist_directory=persist_directory,embedding_function=embeddings)
            print("ChromaDB loaded.")
        return db, embeddings
    except Exception as e:
        print(f"Error initializing ChromaDB: {e}")
        return None, None

def setup_llm():
    try:
        groq_client = Groq(api_key=os.getenv("GROQ_API_KEY"))
        llm = groq_client.chat.completions.create(model="compound-beta-mini",messages=[{"role": "user", "content": "Hello"}])
        print("Connected to Groq.")
        return lambda prompt: groq_client.chat.completions.create(model="llama2",messages=[{"role": "user", "content": prompt}]).choices[0].message.content
    except Exception as e:
        print(f"Could not connect to Groq: {e}. Falling back to simulated LLM.")
        return lambda prompt: f"Simulated LLM response based on:\nQuery: {prompt.split('Context:')[0].strip()}\nContext: {prompt.split('Context:')[1].strip()}"

if __name__ == "__main__":
    os.makedirs(CHROMA_DB_DIR, exist_ok=True)
    chunks = load_and_chunk_docx(DOC_PATH)
    if not chunks:
        print("Exiting due to document loading error.")
        exit()
    vectorstore, embeddings = initialize_chroma_db(chunks, CHROMA_DB_DIR)
    if vectorstore is None:
        print("Exiting due to ChromaDB initialization error.")
        exit()
    retriever = vectorstore.as_retriever()
    llm = setup_llm()

    
    template = """You are an AI assistant. Use the following pieces of context to answer the question at the end.If you don't know the answer, just say that you don't know, don't try to make up an answer.Keep the answer concise and relevant to the provided context.
    Context: {context}
    Question: {question}
    Answer:"""
    prompt = PromptTemplate.from_template(template)
    rag_chain = ({"context": retriever, "question": RunnablePassthrough()} | prompt | llm | StrOutputParser() )
    print("\n--- RAG System Ready ---")
    print("Ask questions about the document. Type 'exit' to quit.")
    while True:
        query = input("\nYour Question: ")
        if query.lower() == 'exit':
            break
        
        # Debugging step: Check if retriever is fetching document chunks
        retrieved_docs = retriever.invoke(query)
        print(f"Retrieved Docs: {retrieved_docs}")  # Add this line
        query = str(query) 
        print("\nSearching and generating response...")
        try:
            # Directly test LLM response before RAG chain
            llm_response = groq_client.chat.completions.create(
            model="compound-beta-mini",
            messages=[{"role": "user", "content": query}])
            print(f"AI (LLM Only): {llm_response.choices[0].message.content}")
            # Now invoke the RAG chain
            response = rag_chain.invoke(query)
            print(f"AI (RAG Response): {response}")
        except Exception as e:
            print(f"An error occurred during response generation: {e}")
        
import os
from docx import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.embeddings import SentenceTransformerEmbeddings
from langchain_community.vectorstores import Chroma
from groq import Groq

# Configuration
DOC_PATH = "C:/Users/admin/Documents/VS CODE/RAG/word.docx"
CHROMA_DB_DIR = "chroma_db_word_rag"
EMBEDDING_MODEL_NAME = "all-MiniLM-L6-v2"

def load_and_chunk_docx(file_path):
    try:
        doc = Document(file_path)
        full_text = [paragraph.text for paragraph in doc.paragraphs]
        text = "\n".join(full_text)
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000,chunk_overlap=200,length_function=len,is_separator_regex=False)
        chunks = text_splitter.create_documents([text])
        print(f"Loaded {len(chunks)} chunks from '{file_path}'")
        return chunks
    except Exception as e:
        print(f"Error loading or chunking document: {e}")
        return []

def initialize_chroma_db(chunks, persist_directory):
    try:
        embeddings = SentenceTransformerEmbeddings(model_name=EMBEDDING_MODEL_NAME)
        print(f"Loading embedding model: {EMBEDDING_MODEL_NAME}...")
        if not os.path.exists(persist_directory):
            print(f"Creating new ChromaDB at '{persist_directory}'...")
            db = Chroma.from_documents(documents=chunks, embedding=embeddings,persist_directory=persist_directory)
            db.persist()
            print("ChromaDB created and persisted.")
        else:
            print(f"Loading existing ChromaDB from '{persist_directory}'...")
            db = Chroma( persist_directory=persist_directory, embedding_function=embeddings )
            print("ChromaDB loaded.")
        return db, embeddings
    except Exception as e:
        print(f"Error initializing ChromaDB: {e}")
        return None, None

def setup_llm():
    try:
        groq_client = Groq(api_key=os.getenv("GROQ_API_KEY"))
        llm = groq_client.chat.completions.create( model="llama2", messages=[{"role": "user", "content": "Hello"}] )
        print("Connected to Groq.")
        return lambda prompt: groq_client.chat.completions.create( model="llama2", messages=[{"role": "user", "content": prompt}] ).choices[0].message.content
    except Exception as e:
        print(f"Could not connect to Groq: {e}. Falling back to simulated LLM.")
        return lambda prompt: f"Simulated LLM response based on:\nQuery: {prompt.split('Context:')[0].strip()}\nContext: {prompt.split('Context:')[1].strip()}"

if __name__ == "__main__":
    os.makedirs(CHROMA_DB_DIR, exist_ok=True)
    chunks = load_and_chunk_docx(DOC_PATH)
    if not chunks:
        print("Exiting due to document loading error.")
        exit()
        vectorstore, embeddings = initialize_chroma_db(chunks, CHROMA_DB_DIR)
        if vectorstore is None:
            print("Exiting due to ChromaDB initialization error.")
            exit()
            retriever = vectorstore.as_retriever()
            llm = setup_llm()
            template = """You are an AI assistant. Use the following pieces of context to answer the question at the end.If you don't know the answer, just say that you don't know, don't try to make up an answer.Keep the answer concise and relevant to the provided context.
            Context: {context}
            Question: {question}
            Answer:"""
            prompt = PromptTemplate.from_template(template)
            rag_chain = ( {"context": retriever, "question": RunnablePassthrough()} | prompt | llm | StrOutputParser() )
            print("\n--- RAG System Ready ---")
            print("Ask questions about the document. Type 'exit' to quit.")
            while True:
                query = input("\nYour Question: ")
                if query.lower() == 'exit':
                    break
                print("\nSearching and generating response...")
                try:
                    response = rag_chain.invoke(query)
                    print(f"AI: {response}")
                except Exception as e:
                    print(f"An error occurred during response generation: {e}") 
                    print("\nSearching and generating response...")
                try:
                    response = rag_chain.invoke(query)
                    print(f"AI: {response}")
                except Exception as e:
                    print(f"An error occurred during response generation: {e}")
                        
                    